import io

import pytest
from docx import Document

from app import export, parsing

RESUME_LINES = [
    "Priya Sharma",
    "priya@example.com | +91 98765 43210",
    "Experience",
    "Software Engineer, Acme Corp (2022 - Present)",
    "Built REST services in Flask serving 2M requests per day",
    "Reduced page load time by 40% by optimizing the React bundle",
    "Education: B.Tech Computer Science, IIT Delhi",
]


def make_docx(lines: list[str], table_row: list[str] | None = None) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if table_row:
        table = doc.add_table(rows=1, cols=len(table_row))
        for cell, value in zip(table.rows[0].cells, table_row):
            cell.text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_text_includes_paragraphs_and_tables():
    data = make_docx(RESUME_LINES, table_row=["Skills", "Python, SQL"])
    text = parsing.extract_text("cv.docx", data)
    assert "Built REST services in Flask" in text
    assert "Skills | Python, SQL" in text


def test_pdf_round_trip(sample_resume):
    text = parsing.extract_text("cv.pdf", export.to_pdf(sample_resume))
    assert "Priya Sharma" in text
    assert "Acme Corp" in text


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("cv.txt", b"hello", "Unsupported file type"),
        ("cv.pdf", b"not a pdf at all", "isn't a valid PDF"),
        ("cv.docx", b"%PDF-1.4 nope", "isn't a valid Word document"),
        ("cv.pdf", b"%PDF-1.4 garbage", "corrupted"),
    ],
)
def test_invalid_files_raise_friendly_errors(filename, data, message):
    with pytest.raises(parsing.ResumeParseError, match=message):
        parsing.extract_text(filename, data)


def test_too_little_text_is_rejected():
    with pytest.raises(parsing.ResumeParseError, match="couldn't read enough text"):
        parsing.extract_text("cv.docx", make_docx(["Hi"]))


def test_normalize_text_unifies_bullets_and_whitespace():
    raw = "  Built   things\r\n\r\n\r\n\r\n● Shipped stuff"
    assert parsing.normalize_text(raw) == "• Built things\n\n• Shipped stuff"
