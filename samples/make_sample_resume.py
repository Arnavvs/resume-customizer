"""Generates sample_resume.docx (and a PDF via the backend exporter is not used on purpose:
the DOCX uses Word bullets and a skills table, like real-world templates)."""
from docx import Document
from docx.shared import Pt

doc = Document()
doc.styles["Normal"].font.size = Pt(10.5)
h = doc.add_paragraph(); r = h.add_run("Rahul Verma"); r.bold = True; r.font.size = Pt(18)
doc.add_paragraph("Software Developer")
doc.add_paragraph("rahul.verma@example.com | +91 98100 12345 | Pune, India | github.com/rahulv")

doc.add_heading("Objective", level=2)
doc.add_paragraph("Seeking a challenging role in a reputed organization where I can utilize my skills "
                  "and grow as a software developer. Hard-working team player with good communication skills.")

doc.add_heading("Technical Skills", level=2)
t = doc.add_table(rows=3, cols=2)
for row, (k, v) in zip(t.rows, [("Languages", "Python, Java, JavaScript, SQL"),
                                ("Frameworks", "Django, Flask, ReactJS, Celery"),
                                ("Tools", "Git, Docker, Postgres, Redis, Jenkins, Linux, AWS S3")]):
    row.cells[0].text, row.cells[1].text = k, v

doc.add_heading("Work Experience", level=2)
p = doc.add_paragraph(); p.add_run("Software Developer — PayKart Technologies").bold = True
doc.add_paragraph("Pune | Jul 2022 – Present")
for b in ["Worked on the frontend dashboard using ReactJS for merchant reports",
          "Developed REST APIs in Django for the payments reconciliation module handling 1.5M transactions per day",
          "Wrote unit tests using pytest and improved coverage from 45% to 80%",
          "Optimized slow Postgres queries, reducing report generation time by 60%",
          "Set up Jenkins pipelines to build Docker images and deploy to staging",
          "Used Celery and Redis for background jobs like sending settlement emails"]:
    doc.add_paragraph(b, style="List Bullet")

p = doc.add_paragraph(); p.add_run("Software Engineering Intern — DataNest Solutions").bold = True
doc.add_paragraph("Remote | Jan 2022 – Jun 2022")
for b in ["Built a Flask microservice to ingest CSV files into PostgreSQL",
          "Helped migrate file storage to AWS S3",
          "Fixed bugs reported by QA team"]:
    doc.add_paragraph(b, style="List Bullet")

doc.add_heading("Projects", level=2)
p = doc.add_paragraph(); p.add_run("ExpenseSplit").bold = True; p.add_run(" — Django, React, PostgreSQL")
for b in ["Made a bill splitting web app used by 300+ students in college",
          "Implemented JWT authentication and role-based access"]:
    doc.add_paragraph(b, style="List Bullet")
p = doc.add_paragraph(); p.add_run("Stock Alert Bot").bold = True; p.add_run(" — Python, Telegram API")
doc.add_paragraph("Telegram bot that sends price alerts for NSE stocks using a scheduled scraper", style="List Bullet")

doc.add_heading("Education", level=2)
doc.add_paragraph("B.E. Computer Engineering, Pune Institute of Computer Technology, 2018 – 2022, CGPA 8.4/10")
doc.add_heading("Certifications", level=2)
doc.add_paragraph("AWS Certified Cloud Practitioner (2023)", style="List Bullet")
doc.save("sample_resume.docx")
print("wrote sample_resume.docx")
